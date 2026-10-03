#!/usr/bin/env python3
"""
Shop Mulmul – Abandoned Carts report (same cut as the Sept 2026 manual reports).

Usage:
  SHOPIFY_STORE=shopmulmul.myshopify.com SHOPIFY_TOKEN=shpat_xxx \
  python abandoned_carts_report.py --period daily|weekly|monthly [--out reports/]
  python abandoned_carts_report.py --from 2026-09-16 --to 2026-09-30

Period semantics (IST):
  daily   -> trailing 7 days ending yesterday (yesterday's row is "fresh", older rows have matured)
  weekly  -> previous Mon–Sun
  monthly -> previous calendar month
Conversions are counted twice: "period" = up to period end 23:59 IST (comparable across reports),
"to-date" = up to run time (addendum columns).

Needs: requests, pandas, openpyxl. Admin API scopes: read_orders, read_customers, read_checkouts.
"""
import argparse, os, re, sys, time, json
from datetime import datetime, timedelta, date
import requests, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

IST = timedelta(hours=5, minutes=30)
API_VER = "2026-01"
STORE = os.environ.get("SHOPIFY_STORE", "shopmulmul.myshopify.com")
TOKEN = os.environ.get("SHOPIFY_TOKEN")

Q = """
query($first:Int!,$after:String,$q:String,$oq:String){
  abandonedCheckouts(first:$first, after:$after, query:$q, sortKey:CREATED_AT){
    pageInfo{hasNextPage endCursor}
    nodes{ createdAt totalPriceSet{shopMoney{amount}}
      shippingAddress{phone} billingAddress{phone}
      customer{ id phone
        orders(first:5, sortKey:CREATED_AT, query:$oq){
          nodes{ name createdAt cancelledAt sourceName note currentTotalPriceSet{shopMoney{amount}} } } } } } }
"""

def gql(query, variables):
    for attempt in range(6):
        r = requests.post(f"https://{STORE}/admin/api/{API_VER}/graphql.json",
                          headers={"X-Shopify-Access-Token": TOKEN, "Content-Type": "application/json"},
                          json={"query": query, "variables": variables}, timeout=60)
        if r.status_code == 429 or (r.ok and "THROTTLED" in r.text):
            time.sleep(2 ** attempt); continue
        r.raise_for_status()
        j = r.json()
        if "errors" in j and not j.get("data"):
            raise RuntimeError(j["errors"])
        return j["data"]
    raise RuntimeError("Shopify throttled repeatedly")

def period_bounds(period, run_ist):
    today = run_ist.date()
    if period == "daily":
        end = today - timedelta(days=1); start = end - timedelta(days=6)
    elif period == "weekly":
        end = today - timedelta(days=today.weekday() + 1); start = end - timedelta(days=6)
    elif period == "monthly":
        first_this = today.replace(day=1); end = first_this - timedelta(days=1); start = end.replace(day=1)
    else:
        raise ValueError(period)
    return start, end

def fetch(start_ist: date, end_ist: date):
    s_utc = datetime.combine(start_ist, datetime.min.time()) - IST
    e_utc = datetime.combine(end_ist + timedelta(days=1), datetime.min.time()) - IST
    q = f"created_at:>='{s_utc:%Y-%m-%dT%H:%M:%S}Z' AND created_at:<'{e_utc:%Y-%m-%dT%H:%M:%S}Z'"
    oq = f"created_at:>='{s_utc:%Y-%m-%dT%H:%M:%S}Z'"
    rows, after = [], None
    while True:
        d = gql(Q, {"first": 50, "after": after, "q": q, "oq": oq})["abandonedCheckouts"]
        for n in d["nodes"]:
            c = n.get("customer") or {}
            phones = [(n.get("shippingAddress") or {}).get("phone"), (n.get("billingAddress") or {}).get("phone"), c.get("phone")]
            phone = next((p for p in phones if p), "")
            orders = [dict(name=o["name"], at=datetime.strptime(o["createdAt"], "%Y-%m-%dT%H:%M:%SZ"),
                           canc=o["cancelledAt"] is not None,
                           src={"web": "web", "944701441": "app", "shopify_draft_order": "draft"}.get(o["sourceName"], "other"),
                           amt=float(o["currentTotalPriceSet"]["shopMoney"]["amount"]), note=o.get("note") or "")
                      for o in ((c.get("orders") or {}).get("nodes") or [])]
            cr = datetime.strptime(n["createdAt"], "%Y-%m-%dT%H:%M:%SZ")
            rows.append(dict(cr=cr, day=(cr + IST).date(), amt=float(n["totalPriceSet"]["shopMoney"]["amount"]),
                             phone=phone, cust=c.get("id") or n["createdAt"], orders=orders))
        if not d["pageInfo"]["hasNextPage"]: break
        after = d["pageInfo"]["endCursor"]
    return rows

def dummy(p):
    if not p: return True
    intl = p.startswith("+") and not p.startswith("+91")
    d = re.sub(r"\D", "", p)
    if intl: return len(d) < 9
    if d.startswith("0091"): d = d[4:]
    if d.startswith("91") and len(d) == 12: d = d[2:]
    d = d.lstrip("0") if len(d) > 10 else (d[1:] if d.startswith("0") else d)
    return len(d) != 10 or len(set(d)) <= 2 or d in ("9876543210", "1234567890") or d[0] not in "6789"

def classify(r, cutoff):
    os_ = sorted([o for o in r["orders"] if r["cr"] <= o["at"] <= cutoff], key=lambda o: o["at"])
    if not os_: return None
    def bucket(o):
        if o["src"] == "draft": return "retail" if "REC AT" in o["note"].upper() else "stylist"
        return {"web": "online", "app": "app"}.get(o["src"], "other")
    first = os_[0]
    if not first["canc"]: return dict(bucket=bucket(first), amt=first["amt"])
    for o in os_[1:]:
        if not o["canc"]:
            if o["src"] == "draft" and (o["at"] - first["at"]) <= timedelta(days=3):
                return dict(bucket={"web": "online_to_draft", "app": "app_to_draft"}.get(first["src"], "other_to_draft"), amt=o["amt"])
            return dict(bucket=bucket(o), amt=o["amt"])
    return dict(bucket="cancelled", amt=0)

BUCKETS = ["online", "app", "stylist", "online_to_draft", "app_to_draft", "other_to_draft", "retail", "cancelled", "other"]
CONV = BUCKETS[:6]

def daily_table(df, col, days):
    out = []
    for day in days:
        g = df[df.day == day]
        r = dict(day=day, total=len(g), abv=g.amt.sum(), fake=g[g.fake].amt.sum()); r["act"] = r["abv"] - r["fake"]
        for b in BUCKETS:
            sel = [c for c in g[col] if c and c["bucket"] == b]
            r[b] = len(sel); r[b + "_v"] = sum(c["amt"] for c in sel)
        out.append(r)
    d = pd.DataFrame(out)
    d["conv_n"] = d[CONV].sum(axis=1); d["conv_v"] = d[[b + "_v" for b in CONV]].sum(axis=1)
    return d

def build_xlsx(dp, dt, start, end, path, run_ist):
    wb = Workbook(); INR = "#,##0"; PCT = "0.0%"
    thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
    FILL = PatternFill("solid", fgColor="F2F2F2")
    def w(ws, r, c, v, blue=False, fmt=None, bold=False, center=False):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name="Arial", size=10, bold=bold, color=("0000FF" if blue else "000000"))
        if fmt: cell.number_format = fmt
        if center: cell.alignment = Alignment(horizontal="center")
        cell.border = BOX; return cell
    lab = lambda d: f"{d.day}{'st' if d.day in (1,21,31) else 'nd' if d.day in (2,22) else 'rd' if d.day in (3,23) else 'th'}"
    title = f"{start:%d %b} – {end:%d %b %Y}"
    # Summary
    ws = wb.active; ws.title = "Summary"
    ws["A1"] = f"ABANDONED CARTS ({title})"; ws["A1"].font = Font(name="Arial", bold=True, size=12)
    ws.merge_cells("G2:J2"); ws["G2"] = "Customers converted"; ws.merge_cells("K2:N2"); ws["K2"] = "Converted value (₹)"
    for c in ("G2", "K2"): ws[c].font = Font(name="Arial", bold=True, size=11); ws[c].alignment = Alignment(horizontal="center")
    for i, h in enumerate(["", "total", "Abandoned value", "Fake/ no number", "fake %", "actual abandonment", "Online", "App", "Stylist", "Total", "Online", "App", "Stylist", "total converted", "conversion %"], 1):
        w(ws, 3, i, h, bold=True, center=True).fill = FILL
    r = 4
    for _, d in dp.iterrows():
        w(ws, r, 1, lab(d.day), bold=True); w(ws, r, 2, int(d.total), True); w(ws, r, 3, round(d.abv), True, INR); w(ws, r, 4, round(d.fake), True, INR)
        w(ws, r, 5, f"=IFERROR(D{r}/C{r},0)", fmt=PCT); w(ws, r, 6, f"=C{r}-D{r}", fmt=INR)
        w(ws, r, 7, int(d.online + d.online_to_draft), True); w(ws, r, 8, int(d.app + d.app_to_draft), True); w(ws, r, 9, int(d.stylist + d.other_to_draft), True)
        w(ws, r, 10, f"=G{r}+H{r}+I{r}")
        w(ws, r, 11, round(d.online_v + d.online_to_draft_v), True, INR); w(ws, r, 12, round(d.app_v + d.app_to_draft_v), True, INR); w(ws, r, 13, round(d.stylist_v + d.other_to_draft_v), True, INR)
        w(ws, r, 14, f"=K{r}+L{r}+M{r}", fmt=INR); w(ws, r, 15, f"=IFERROR(N{r}/F{r},0)", fmt=PCT); r += 1
    last = r - 1; w(ws, r, 1, "Period", bold=True)
    for c in [2, 3, 4, 6, 7, 8, 9, 10, 11, 12, 13, 14]:
        L = get_column_letter(c); w(ws, r, c, f"=SUM({L}4:{L}{last})", fmt=(INR if c >= 3 else None), bold=True)
    w(ws, r, 5, f"=IFERROR(D{r}/C{r},0)", fmt=PCT, bold=True); w(ws, r, 15, f"=IFERROR(N{r}/F{r},0)", fmt=PCT, bold=True)
    notes = [f"How this was built (Shopify Admin API, pulled {run_ist:%d %b %Y %H:%M} IST; blue = pulled values, black = formulas):",
             "• Day = abandoned-checkout date in IST; one row per customer per day (highest-value cart kept). Fake/ no number = no phone anywhere, or a dummy number.",
             f"• Converted = the customer's first order after abandoning, up to {end:%d %b} 11:59 PM IST, at current order value. Cancelled orders excluded. Later conversions are on 'By channel' as to-date addendum columns.",
             "• Channel: Online = Online Store, App = Appbrew, Stylist = draft orders (REC AT drafts excluded, re-books split out on 'By channel').",
             "• Only the first 5 orders per customer since period start were read."]
    for i, n in enumerate(notes): ws.cell(row=r + 2 + i, column=1, value=n).font = Font(name="Arial", size=9, italic=(i > 0))
    for c, wd in zip("ABCDEFGHIJKLMNO", [11, 7, 16, 15, 8, 18, 8, 7, 8, 7, 11, 10, 11, 15, 12]): ws.column_dimensions[c].width = wd
    ws.freeze_panes = "B4"
    # By channel
    ws = wb.create_sheet("By channel")
    ws["A1"] = f"CONVERTED ABANDONERS BY CHANNEL BUCKET ({title})"; ws["A1"].font = Font(name="Arial", bold=True, size=12)
    ws.merge_cells("C2:I2"); ws["C2"] = "Customers converted"; ws.merge_cells("J2:P2"); ws["J2"] = "Converted value (₹)"
    for c in ("C2", "J2"): ws[c].font = Font(name="Arial", bold=True, size=11); ws[c].alignment = Alignment(horizontal="center")
    heads = ["", "actual abandonment", "online", "app", "stylist", "online_to_draft", "app_to_draft", "other_to_draft", "Total", "online", "app", "stylist", "online_to_draft", "app_to_draft", "other_to_draft", "Total", "conversion %", "cancelled after converting", "retail_assist (REC AT)", "to-date extra conversions", "to-date extra value (₹)"]
    for i, h in enumerate(heads, 1): w(ws, 3, i, h, bold=True, center=True).fill = FILL
    r = 4
    for (_, d), (_, t) in zip(dp.iterrows(), dt.iterrows()):
        w(ws, r, 1, lab(d.day), bold=True); w(ws, r, 2, round(d.act), True, INR)
        for i, b in enumerate(CONV): w(ws, r, 3 + i, int(d[b]), True); w(ws, r, 10 + i, round(d[b + "_v"]), True, INR)
        w(ws, r, 9, f"=SUM(C{r}:H{r})"); w(ws, r, 16, f"=SUM(J{r}:O{r})", fmt=INR); w(ws, r, 17, f"=IFERROR(P{r}/B{r},0)", fmt=PCT)
        w(ws, r, 18, int(d.cancelled), True); w(ws, r, 19, int(d.retail), True)
        w(ws, r, 20, int(t.conv_n - d.conv_n), True); w(ws, r, 21, round(t.conv_v - d.conv_v), True, INR); r += 1
    last = r - 1; w(ws, r, 1, "Period", bold=True)
    for c in list(range(2, 17)) + [18, 19, 20, 21]:
        L = get_column_letter(c); w(ws, r, c, f"=SUM({L}4:{L}{last})", fmt=(INR if c in (2, 10, 11, 12, 13, 14, 15, 16, 21) else None), bold=True)
    w(ws, r, 17, f"=IFERROR(P{r}/B{r},0)", fmt=PCT, bold=True)
    for c, wd in zip([get_column_letter(i) for i in range(1, 22)], [11, 17, 8, 7, 8, 14, 12, 13, 7, 11, 10, 11, 14, 12, 13, 11, 12, 14, 14, 16, 14]): ws.column_dimensions[c].width = wd
    ws.freeze_panes = "B4"
    wb.save(path)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--period", choices=["daily", "weekly", "monthly"])
    ap.add_argument("--from", dest="from_"); ap.add_argument("--to"); ap.add_argument("--out", default="reports")
    a = ap.parse_args()
    if not TOKEN: sys.exit("SHOPIFY_TOKEN missing")
    run_ist = datetime.utcnow() + IST
    if a.period: start, end = period_bounds(a.period, run_ist); tag = a.period
    else: start, end = date.fromisoformat(a.from_), date.fromisoformat(a.to); tag = "custom"
    rows = fetch(start, end)
    df = pd.DataFrame(rows)
    if df.empty: sys.exit("no abandoned checkouts in range")
    df["fake"] = df.phone.apply(dummy)
    df = df.sort_values("amt", ascending=False).drop_duplicates(["cust", "day"]).sort_values("cr")
    cut_period = datetime.combine(end + timedelta(days=1), datetime.min.time()) - IST
    df["period"] = df.apply(lambda r: classify(r, cut_period), axis=1)
    df["todate"] = df.apply(lambda r: classify(r, datetime.utcnow()), axis=1)
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    dp, dt = daily_table(df, "period", days), daily_table(df, "todate", days)
    os.makedirs(a.out, exist_ok=True)
    path = os.path.join(a.out, f"abandoned_carts_{tag}_{start:%Y-%m-%d}_{end:%Y-%m-%d}.xlsx")
    build_xlsx(dp, dt, start, end, path, run_ist)
    t, tt = dp.sum(numeric_only=True), dt.sum(numeric_only=True)
    summary = {
        "period": f"{start} to {end}", "carts": int(t.total), "abandoned_value": round(t.abv), "fake_value": round(t.fake),
        "fake_pct": round(t.fake / t.abv, 4) if t.abv else 0, "actual_abandonment": round(t.act),
        "converted_customers": int(t.conv_n), "converted_value": round(t.conv_v),
        "recovery_pct": round(t.conv_v / t.act, 4) if t.act else 0,
        "to_date_converted_customers": int(tt.conv_n), "to_date_converted_value": round(tt.conv_v),
        "to_date_recovery_pct": round(tt.conv_v / t.act, 4) if t.act else 0,
        "channel_value": {b: round(t[b + "_v"]) for b in CONV + ["retail"]},
        "cancelled_after_converting": int(t.cancelled), "file": path,
    }
    # yesterday callout for daily runs
    if a.period == "daily":
        y = dt[dt.day == end].iloc[0]
        summary["yesterday"] = {"carts": int(y.total), "abandoned_value": round(y.abv), "actual": round(y.act), "converted_so_far": int(y.conv_n), "value_so_far": round(y.conv_v)}
    with open(os.path.join(a.out, f"abandoned_carts_{tag}_latest.json"), "w") as f: json.dump(summary, f, indent=2, default=str)
    print(json.dumps(summary, indent=2, default=str))

if __name__ == "__main__":
    main()
